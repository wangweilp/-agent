"""DOCX (.docx) file parser using python-docx."""

import os
import re
from datetime import datetime, timezone

from src.parsers.base import DocumentChunk


def _extract_tags(text: str, max_terms: int = 5) -> list[str]:
    """Extract simple tags from text."""
    terms: list[str] = []
    code_terms = set(re.findall(r"\b[A-Z][a-z]+(?:[A-Z][a-z]+)+\b", text))
    terms.extend(code_terms)
    bracket_terms = re.findall(r"\[([^\]]+)\]", text)
    terms.extend(bracket_terms)
    bold_terms = re.findall(r"\*\*([^*]+)\*\*", text)
    terms.extend(bold_terms)
    seen: set[str] = set()
    unique = []
    for t in terms:
        t = t.strip().lower()
        if t and t not in seen and len(t) > 2:
            seen.add(t)
            unique.append(t)
    return unique[:max_terms]


class DocxParser:
    """Parse .docx files into chunks by paragraph groups."""

    def parse(self, file_path: str) -> list[DocumentChunk]:
        """Parse a .docx file. Never raises — returns empty list on error."""
        try:
            if not os.path.isfile(file_path):
                return []

            text = self._extract_text(file_path)
            if not text.strip():
                return []

            source_basename = os.path.basename(file_path)

            # Split by paragraph breaks
            paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

            from src.parsers.chunking import paragraph_chunks

            chunk_texts = paragraph_chunks(paragraphs, chunk_size=2000, overlap=250)

            chunks: list[DocumentChunk] = []
            for idx, chunk_text in enumerate(chunk_texts):
                title = chunk_text[:200].strip()
                tags = _extract_tags(chunk_text)
                chunks.append(
                    DocumentChunk(
                        id=f"{source_basename}_chunk_{idx}",
                        source=file_path,
                        title=title,
                        content=chunk_text,
                        tags=tags,
                        created_at=datetime.now(timezone.utc),
                        metadata={
                            "char_count": len(chunk_text),
                            "source_file": file_path,
                            "parser": "docx",
                            "chunk_index": idx,
                        },
                    )
                )

            return chunks

        except Exception:
            return []

    @staticmethod
    def _extract_text(file_path: str) -> str:
        """Extract all text from a .docx file, preserving paragraph structure."""
        try:
            from docx import Document

            doc = Document(file_path)
            paragraphs: list[str] = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    # Detect heading style
                    if para.style and para.style.name and para.style.name.startswith("Heading"):
                        paragraphs.append(f"## {text}")
                    else:
                        paragraphs.append(text)

            # Also extract table text
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells]
                    paragraphs.append(" | ".join(row_texts))

            return "\n\n".join(paragraphs)
        except Exception:
            return ""
