"""DOCX file importer — extracts text via python-docx."""

from __future__ import annotations

from pathlib import Path

from src.importers.base import ImportResult, chunk_text, extract_entities, generate_title


class DocxImporter:
    """Import .docx files, extracting paragraph text."""

    def import_data(self, source: str) -> ImportResult:
        path = Path(source)
        metadata: dict = {"file_path": str(path.absolute()), "format": "docx"}

        if not path.is_file():
            return ImportResult(
                source="docx",
                title=path.stem,
                error=f"File not found: {source}",
                metadata=metadata,
            )

        try:
            import docx  # type: ignore[import-untyped]
        except ImportError:
            return ImportResult(
                source="docx",
                title=path.stem,
                error="python-docx is not installed",
                metadata=metadata,
            )

        try:
            doc = docx.Document(str(path))
        except Exception as exc:
            return ImportResult(
                source="docx",
                title=path.stem,
                error=f"Failed to open document: {exc}",
                metadata=metadata,
            )

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        if not paragraphs:
            # Try extracting from tables as well
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        paragraphs.append(" | ".join(row_texts))

        text = "\n\n".join(paragraphs)

        if not text.strip():
            return ImportResult(
                source="docx",
                title=path.stem,
                chunks=[],
                entities=[],
                metadata=metadata,
            )

        title = generate_title(text, fallback=path.stem)
        entities = extract_entities(text)
        text_chunks = chunk_text(text)

        chunks: list[dict] = []
        for text_segment in text_chunks:
            chunks.append({
                "content": text_segment,
                "metadata": {
                    "chunk_index": len(chunks),
                    "source_file": path.name,
                },
            })

        return ImportResult(
            source="docx",
            title=title,
            chunks=chunks,
            entities=entities,
            memory_type="semantic",
            metadata=metadata,
        )
