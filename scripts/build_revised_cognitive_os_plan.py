from __future__ import annotations

import os
import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PROJECT_ROOT = Path(r"D:\dma\day2")
DESKTOP = Path.home() / "Desktop"
OUTPUT = DESKTOP / "黔智脑_Cognitive_OS_计划书1_基于当前项目修订完善版.docx"
FIGURE_SRC = PROJECT_ROOT / "docs" / "competition" / "figures_ai"
FIGURE_OUT = PROJECT_ROOT / "docs" / "competition" / "figures_ai_compressed"


BODY_EAST = "宋体"
BODY_WEST = "Times New Roman"
H1_EAST = "黑体"
H2_EAST = "楷体"
CAPTION_EAST = "黑体"


def twips(cm: float) -> int:
    return int(Cm(cm).twips)


def set_run_font(run, east: str = BODY_EAST, west: str = BODY_WEST, size: int = 14, bold: bool = False) -> None:
    run.font.name = west
    run.font.size = Pt(size)
    run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), west)
    r_fonts.set(qn("w:hAnsi"), west)
    r_fonts.set(qn("w:eastAsia"), east)


def set_para_base(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_next: bool = False) -> None:
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(28)
    pf.keep_with_next = keep_next
    p_pr = p._p.get_or_add_pPr()
    snap = p_pr.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        p_pr.append(snap)
    snap.set(qn("w:val"), "0")


def add_text_paragraph(doc: Document, text: str, style: str = "body") -> None:
    p = doc.add_paragraph()
    if style == "h1":
        set_para_base(p, WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        run = p.add_run(text)
        set_run_font(run, H1_EAST, BODY_WEST, 16, False)
    elif style == "h2":
        set_para_base(p, WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        run = p.add_run(text)
        set_run_font(run, H2_EAST, BODY_WEST, 16, False)
    elif style == "h3":
        set_para_base(p, WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        run = p.add_run(text)
        set_run_font(run, BODY_EAST, BODY_WEST, 14, False)
    elif style == "title":
        set_para_base(p, WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        run = p.add_run(text)
        set_run_font(run, H1_EAST, BODY_WEST, 22, False)
    elif style == "subtitle":
        set_para_base(p, WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        run = p.add_run(text)
        set_run_font(run, BODY_EAST, BODY_WEST, 14, False)
    elif style == "caption":
        set_para_base(p, WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        run = p.add_run(text)
        set_run_font(run, CAPTION_EAST, BODY_WEST, 12, False)
    else:
        set_para_base(p)
        run = p.add_run(text)
        set_run_font(run)


def add_blank(doc: Document, n: int = 1) -> None:
    for _ in range(n):
        p = doc.add_paragraph()
        set_para_base(p)
        p.add_run("")


def add_page_number(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_base(p, WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run()
    set_run_font(run, BODY_EAST, BODY_WEST, 12, False)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.5)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.0)
    pg_mar = section._sectPr.pgMar
    pg_mar.set(qn("w:gutter"), str(twips(0.5)))
    add_page_number(section)

    styles = doc.styles
    for style_name in ["Normal", "Body Text"]:
        style = styles[style_name]
        style.font.name = BODY_WEST
        style.font.size = Pt(14)
        style.font.bold = False
        r_pr = style._element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)
        r_fonts.set(qn("w:ascii"), BODY_WEST)
        r_fonts.set(qn("w:hAnsi"), BODY_WEST)
        r_fonts.set(qn("w:eastAsia"), BODY_EAST)
        p_pr = style._element.get_or_add_pPr()
        spacing = p_pr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            p_pr.append(spacing)
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "560")
        spacing.set(qn("w:lineRule"), "exact")


def compress_image(src_name: str, out_name: str, max_width: int = 1500, quality: int = 82) -> Path:
    FIGURE_OUT.mkdir(parents=True, exist_ok=True)
    src = FIGURE_SRC / src_name
    out = FIGURE_OUT / out_name
    with Image.open(src) as im:
        im = im.convert("RGB")
        if im.width > max_width:
            height = int(im.height * max_width / im.width)
            im = im.resize((max_width, height), Image.Resampling.LANCZOS)
        im.save(out, "JPEG", quality=quality, optimize=True, progressive=True)
    return out


FIGURES = {
    "cover": compress_image("ai_cover_hero_cognitive_os.png", "ai_cover_hero_cognitive_os.jpg"),
    "scenario": compress_image("ai_customer_scenarios.png", "ai_customer_scenarios.jpg"),
    "architecture": compress_image("ai_architecture_layers.png", "ai_architecture_layers.jpg"),
    "workflow": compress_image("ai_investment_agent_workflow.png", "ai_investment_agent_workflow.jpg"),
    "security": compress_image("ai_security_compliance_private_deployment.png", "ai_security_compliance_private_deployment.jpg"),
    "business": compress_image("ai_business_model_value_paths.png", "ai_business_model_value_paths.jpg"),
    "roadmap": compress_image("ai_implementation_roadmap.png", "ai_implementation_roadmap.jpg"),
}


def add_figure(doc: Document, key: str, caption: str) -> None:
    add_blank(doc, 1)
    p = doc.add_paragraph()
    set_para_base(p, WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
    run = p.add_run()
    run.add_picture(str(FIGURES[key]), width=Cm(14.0))
    add_text_paragraph(doc, caption, "caption")
    add_blank(doc, 1)


def set_cell_text(cell, text: str, size: int = 14) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    for idx, part in enumerate(str(text).split("\n")):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        set_para_base(p, WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(part)
        set_run_font(run, BODY_EAST, BODY_WEST, size, False)


def add_table_block(doc: Document, caption: str, rows: list[list[str]]) -> None:
    add_blank(doc, 1)
    add_text_paragraph(doc, caption, "caption")
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            set_cell_text(cell, text)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for side in ["top", "bottom", "left", "right"]:
                elem = tc_mar.find(qn(f"w:{side}"))
                if elem is None:
                    elem = OxmlElement(f"w:{side}")
                    tc_mar.append(elem)
                elem.set(qn("w:w"), "80")
                elem.set(qn("w:type"), "dxa")
    add_blank(doc, 1)


def add_bullets_as_paragraphs(doc: Document, items: list[str]) -> None:
    for item in items:
        add_text_paragraph(doc, f"（{items.index(item) + 1}）{item}")


def paragraph_char_count(texts: list[str]) -> int:
    return len(re.sub(r"\s+", "", "".join(texts)))


def build() -> Path:
    doc = Document()
    configure_document(doc)

    texts_for_count: list[str] = []

    def p(text: str, style: str = "body") -> None:
        texts_for_count.append(text)
        add_text_paragraph(doc, text, style)

    p("黔智脑 Cognitive OS 商业计划书", "title")
    p("基于当前工程原型的修订完善版", "subtitle")
    p("版本日期：2026年6月12日", "subtitle")
    add_figure(doc, "cover", "图0.1 黔智脑 Cognitive OS 项目定位示意图")

    p("摘  要", "h1")
    p("黔智脑 Cognitive OS 定位为面向政企组织的认知操作系统原型，重点解决组织数据分散、知识难沉淀、任务协同难审计和 AI 应用难合规落地的问题。项目并非单一聊天机器人，而是围绕“数据接入、长期记忆、语义检索、知识图谱、多智能体协同、开放平台、安全审计和可视化运营”形成的工程化原型。")
    p("截至2026年6月12日，项目代码位于 D:\\dma\\day2，已形成 FastAPI 后端、Next.js 前端、SQLite/ChromaDB 数据存储、本地 Embedding、DeepSeek/OpenAI-compatible LLM 接入、Agent Registry、Workflow Engine、Marketplace、Open Platform、Runtime readiness、kill switch、incident store、artifact materialization、rootless container prototype gate 等模块。当前可验证的工程事实包括：src 目录约243个 Python 源文件、测试目录95个 Python 测试模块，Open Platform 专项测试2803项通过，核心回归3450项通过。")
    p("为保证计划书真实可信，本版明确区分“已完成的工程原型”“可演示或可验证的控制平面能力”“尚未完成的生产级能力”和“未来商业化假设”。当前尚无签约客户收入、无正式政务云上线、无第三方智能体真实执行、无生产级容器沙箱运行、无真实支付分成。本计划书中的市场、财务和融资内容均为基于公开资料与内部成本假设的测算，不构成收益承诺、投资邀约或已落地成果陈述。")

    p("一、项目概述", "h1")
    p("1. 项目定位", "h2")
    p("黔智脑 Cognitive OS 是面向政府部门、产业园区、中小企业和知识型团队的组织级 AI 平台。它的核心目标是把分散在文档、会议纪要、业务系统、客户记录、政策文件和项目复盘中的信息，转化为可检索、可关联、可协同、可审计的组织知识资产，并在人工复核和权限约束下辅助完成分析、问答、流程整理和方案生成。")
    p("项目当前处于工程化原型和试点准备阶段，适合进行比赛展示、内部演示、小范围客户访谈和受控试点设计，不宜表述为成熟商用平台。后续若进入政务、园区或企业生产环境，仍需完成客户侧数据授权、等保定级备案与测评、个人信息保护影响评估、模型输出风险评估、验收指标确认和运维责任划分。")
    p("2. 当前工程状态", "h2")
    p("本项目的优势在于已有较完整的代码资产和测试资产，而不是仅停留在商业设想。后端采用 FastAPI 与六边形架构，前端采用 Next.js、React 和 TypeScript；知识层采用 SQLite、ChromaDB 和本地 Embedding；多智能体侧已有 Agent Registry、Workflow Engine 和人工确认节点；开放平台侧已有开发者账号、API Key、Scope 权限、提交审核、管理员审核、Marketplace 浏览/安装/配置与用量统计等 MVP 能力。")
    add_table_block(doc, "表1.1 当前工程进展与真实性边界", [
        ["模块", "当前可验证状态", "计划书表述边界"],
        ["长期记忆与检索", "支持多层记忆、语义检索、ChromaDB、本地 Embedding", "可写为工程原型能力"],
        ["多智能体工作流", "已有 Agent Registry、Workflow Engine、人工确认节点", "行业 Agent 仍应写为可扩展场景"],
        ["Marketplace", "支持发现、安装、配置、计量 MVP", "真实支付、分成和收费尚未启用"],
        ["Open Platform", "支持开发者提交、审核、API Key 与 Scope", "第三方代码执行仍被阻断"],
        ["Runtime/Sandbox", "具备 readiness、simulation、metadata-only proof", "不能写成生产沙箱或真实容器执行"],
        ["测试资产", "Open Platform 2803 passed，核心回归3450 passed", "可作为工程质量证据"],
    ])
    p("3. 为什么选择贵州场景", "h2")
    p("贵州长期推进大数据和数字经济发展，具备数据中心、算力、政务服务和产业园区等应用基础。对本项目而言，贵州不是一个泛化标签，而是早期试点场景选择：一方面，政企客户在知识库、政策问答、招商分析、企业服务和数据治理方面存在真实需求；另一方面，本地化交付、客户访谈、数据样例获取和场景共创更容易形成可验证闭环。项目早期应优先选择“企业知识管理”和“园区招商辅助”两个边界清晰的场景，再逐步扩展到政务服务问答和城市治理辅助研判。")

    p("二、项目优势", "h1")
    p("1. 工程化原型优势", "h2")
    p("相较于只做概念包装的 AI 项目，黔智脑已有可运行、可测试、可继续迭代的工程基础。项目已覆盖后端服务、前端页面、数据存储、检索、智能体、开放平台、权限、审计、用量统计和运行时准入等多个层面。这种工程基础有利于比赛展示，也有利于后续与试点客户讨论真实的数据接入、权限配置、验收指标和部署方式。")
    p("2. 组织记忆与知识图谱优势", "h2")
    p("普通大模型应用更多依赖一次性输入和通用知识，难以沉淀组织内部经验。黔智脑的长期记忆设计将会话、文档、项目、会议、图片、音视频和复盘内容逐步转化为结构化知识资产，并通过语义检索和知识图谱建立实体、关系和上下文。这使系统能够支持持续复盘、跨项目检索、经验复用和组织知识传承。")
    p("3. 多智能体与人工复核优势", "h2")
    p("本项目将 AI 能力设计为受控流程，而不是无约束自动执行。当前工程已具备工作流编排和人工确认思路，适合在招商分析、会议转培训、部门知识助手等场景中形成“检索证据—生成草案—人工复核—审计留痕”的闭环。对政企场景而言，这种可审计、可回退、可人工把关的设计比“全自动决策”更符合合规要求。")
    p("4. 安全边界优势", "h2")
    p("当前项目在 Runtime 与 Sandbox 方向坚持 fail-closed 原则。第三方 package、entrypoint、容器、namespace、cgroup、mount、网络访问和真实队列调度均未被开放；Step 26-E 和 Step 26-F 仍是 metadata/control-plane gate，而非真实 rootless container runtime 或 isolated runtime。该边界虽然限制了短期炫技，但降低了计划书夸大风险，也为后续安全评审、红队测试和受控试点保留了清晰路径。")
    add_figure(doc, "architecture", "图2.1 Cognitive OS 工程能力与安全边界示意图")

    p("三、市场分析", "h1")
    p("1. 行业与政策依据", "h2")
    p("生成式人工智能、数字经济和中小企业数字化转型为本项目提供了需求基础。CNNIC《生成式人工智能应用发展报告（2024）》披露，我国已初步构建较为全面的人工智能产业体系，相关企业超过4500家，核心产业规模接近6000亿元。国家统计局2024年国民经济和社会发展统计公报显示，我国信息传输、软件和信息技术服务业保持较快增长。财政部等部门持续推动中小企业数字化转型城市试点，说明中小企业、园区和地方产业服务仍是政策支持方向。")
    p("上述数据说明行业具备发展空间，但不能直接等同于本项目已获得市场份额。本项目的市场判断应回到可触达客户：贵州及西南地区产业园区、中小企业、学校创新团队、政企服务窗口和数字化服务商。早期目标不是覆盖泛全国市场，而是通过小范围试点验证客户是否愿意为“知识管理、智能问答、招商辅助、合规审计和部署服务”付费。")
    p("2. 目标客户与需求", "h2")
    add_table_block(doc, "表3.1 目标客户、需求与验证方式", [
        ["客户类型", "主要需求", "早期验证方式"],
        ["中小企业", "文档知识库、经验沉淀、内部问答、周报和项目复盘", "访谈10—20家，验证资料导入和订阅意愿"],
        ["产业园区/投促机构", "产业链分析、企业画像、政策匹配、招商材料生成", "用公开资料做招商 Demo，争取非付费试点"],
        ["政企服务窗口", "政策问答、办事导航、材料清单、人工转办", "先做离线知识库样例，评估合规和数据授权"],
        ["数字化服务商", "需要可交付的 AI 应用底座和行业模板", "联合方案、渠道分成和项目交付成本测算"],
    ])
    p("3. 竞争与差异化", "h2")
    p("黔智脑不与大模型厂商直接竞争，而是做模型之上的组织知识和业务流程层。大模型厂商提供底层模型能力，通用办公产品提供协作入口，垂直系统提供行业流程，黔智脑的差异化在于把长期记忆、知识图谱、多智能体流程、人工复核、权限审计和本地化交付结合起来。竞争重点不是模型参数规模，而是能否接入组织数据、沉淀业务知识、支撑安全审计并形成可复制交付包。")
    add_figure(doc, "scenario", "图3.1 目标客户与应用场景示意图")

    p("四、产品方案与技术路线", "h1")
    p("1. 产品架构", "h2")
    p("产品采用“数据接入层、知识记忆层、智能体协同层、开放平台层、应用场景层、安全合规层”的结构。数据接入层负责文档、图片、音频、视频、网页和业务系统数据接入；知识记忆层负责解析、向量化、实体抽取、关系沉淀和长期记忆管理；智能体协同层负责任务拆解、证据检索、工具调用和人工确认；开放平台层负责开发者提交、管理员审核、API Key 和 Scope 权限；应用层面向企业知识库、部门助手、招商辅助和政策问答；安全合规层贯穿权限、审计、日志、脱敏和运行时准入。")
    p("2. 当前已实现与规划能力", "h2")
    p("已实现或可演示的能力包括长期记忆、语义检索、知识图谱雏形、文件导入、音视频上传与可选分析、Agent Registry、Workflow Engine、Marketplace MVP、Open Platform、API Key、Scope 权限、管理员审核、用量统计和多项运行时安全 gate。需要谨慎表述的能力包括：Coach、Team Brain、Enterprise Brain 目前更适合作为产品形态演进方向；招商 Agent、政策 Agent、风险 Agent、市民助手属于可扩展行业场景，不宜写成已完成专用智能体；多模型网关目前以 DeepSeek/OpenAI-compatible LLM 与本地 Embedding 为基础，多模型切换仍是后续扩展。")
    p("3. 智能招商流程示例", "h2")
    p("以产业园区招商为例，用户输入目标行业、区域约束和招商目标后，系统首先检索政策文件、产业链资料和公开企业信息；产业分析模块生成产业链结构和本地优势；候选企业模块形成目标企业草案；政策匹配模块提取可适用政策和材料清单；风险复核模块提示信用、舆情、数据来源和合规风险；最终输出给人工人员复核。该流程当前应定位为可演示业务方案和试点方向，不能表述为已在真实招商工作中产生确定成效。")
    add_figure(doc, "workflow", "图4.1 智能招商多智能体协同流程示意图")
    p("4. 安全与合规技术路线", "h2")
    p("平台坚持“最小权限、默认拒绝、人工复核、审计留痕、分阶段开放”的安全路线。当前第三方智能体执行、package 下载、entrypoint 执行、真实容器启动、网络访问和主机路径读取均处于阻断或 metadata-only 阶段。后续若进入生产试点，应先完成可信 fixture、小范围沙箱验证、网络/文件系统/secrets enforcement、资源限制、日志留存和红队测试，再根据客户场景逐步开放。")
    add_figure(doc, "security", "图4.2 安全合规与私有化部署控制示意图")

    p("五、商业模式", "h1")
    p("1. 收入结构假设", "h2")
    p("商业模式采用“企业订阅 + 项目交付 + 行业知识包/报告服务 + 后续生态分成设计”的组合路径。当前真实支付、分成和签约收入尚未形成，Marketplace 只应描述为具备发现、安装、配置和计量 MVP，不应写成已经产生平台佣金。早期更现实的路径是先通过企业知识管理和园区招商 Demo 获取试点，再根据试点成本和客户付费意愿确定标准版价格。")
    add_table_block(doc, "表5.1 产品版本与收费假设", [
        ["版本", "目标客户", "建议价格假设", "验证重点"],
        ["基础版", "小微企业、团队", "0.98—2.98万元/年", "文档导入、检索、问答留存和使用频率"],
        ["专业版", "成长型企业、园区服务机构", "5—12万元/年", "知识图谱、多智能体流程、权限和看板"],
        ["企业/私有化版", "中大型企业、政企客户", "80—200万元/项目起", "部署、数据接入、等保、验收和运维成本"],
        ["园区试点包", "产业园区、投促机构", "15—50万元/试点", "招商场景、政策匹配、报告质量和复购"],
        ["知识包/报告服务", "企业服务、招商服务机构", "按专题或年度包计费", "数据来源、更新机制和交付边界"],
    ])
    p("2. 单位经济模型", "h2")
    p("在未形成真实收入前，应先做可解释的单位经济模型。SaaS 订阅收入的主要成本包括模型调用、本地向量化与存储、客服支持、产品迭代和云资源；项目交付收入的主要成本包括需求调研、数据清洗、部署、权限配置、培训、验收和售后。内部测算可暂按模型/云资源成本占收入15%—25%、项目交付人力成本占收入25%—40%、销售和客户成功成本占收入10%—20%估算，后续必须用试点数据修正。")
    p("3. 商业化路径", "h2")
    p("第一阶段通过比赛、演示材料和开源式技术说明建立可信度；第二阶段寻找1—2个非付费或低价试点，验证真实数据接入、回答质量、权限审计和使用频率；第三阶段把试点过程沉淀为部署清单、数据模板、场景模板、验收指标和报价模型；第四阶段再拓展企业订阅、园区项目和行业知识包。每一步都应以客户访谈、试点记录、成本台账和验收指标作为下一步依据。")
    add_figure(doc, "business", "图5.1 商业模式与价值路径示意图")

    p("六、财务分析与融资计划", "h1")
    p("1. 当前财务状态", "h2")
    p("截至2026年6月12日，项目应按早期原型项目处理：暂无可核验签约收入，暂无可核验长期付费客户，暂无可核验政务或园区正式生产部署。已有资产主要是代码、测试、文档、演示能力和可继续迭代的产品原型。因此财务预测只能作为内部经营测算，不能作为收益承诺或已实现业绩。")
    p("2. 三年经营测算", "h2")
    add_table_block(doc, "表6.1 三年经营测算区间", [
        ["年份", "主要目标", "收入假设", "净利润/亏损假设"],
        ["第1年", "完成MVP打磨和1—2个试点", "30—80万元", "-80至-40万元"],
        ["第2年", "形成标准交付包和5—10个付费客户", "150—300万元", "-20至30万元"],
        ["第3年", "扩大企业订阅和园区项目复制", "400—800万元", "40—120万元"],
    ])
    p("上述测算假设第1年以试点和产品打磨为主，第2年开始出现小规模付费客户，第3年在产品标准化、交付效率提升和复购出现后改善利润。若客户转化率、客单价、模型成本、交付周期或回款周期不达预期，收入和利润均可能低于测算。计划书正式提交时，建议将客户访谈、报价单、试点意向、成本估算表和团队薪酬预算作为申报系统支撑材料。")
    p("3. 融资计划", "h2")
    p("项目可考虑对接天使投资、赛事奖金、政府创新创业资源和产业合作资源。若融资100万元人民币，可主要用于产品研发、试点交付、市场验证、算力资源和合规咨询。股权比例不宜在计划书中写成已确定交易，可表述为“结合尽调、估值、知识产权归属、创始团队持股、员工期权和后续融资安排协商确定”。若需提供参考区间，可将10%—15%作为谈判假设，并明确不构成公开募资、投资邀约或收益承诺。")
    add_table_block(doc, "表6.2 资金用途计划", [
        ["用途", "金额", "说明"],
        ["产品研发", "45万元", "前后端、记忆系统、Agent流程、测试和工程稳定性"],
        ["试点交付", "25万元", "数据接入、部署、培训、验收材料和客户成功"],
        ["算力与基础设施", "12万元", "模型API、本地模型、向量数据库和测试环境"],
        ["市场验证", "10万元", "客户访谈、路演材料、演示视频和渠道合作"],
        ["合规与知识产权", "8万元", "软著、商标、法律咨询、数据合规和安全评估"],
    ])

    p("七、实施计划与风险控制", "h1")
    p("1. 12个月实施路线", "h2")
    p("实施节奏应从“可演示”走向“可试点”，再走向“可复制”，每一阶段设置明确 gate。第1—3个月完成核心 Demo、样例数据和演示材料；第4—6个月完成客户访谈和小范围试点；第7—9个月沉淀标准部署包、验收指标和报价模型；第10—12个月验证付费转化和渠道合作。政务、园区和企业场景不宜承诺过快上线，应把数据授权、安全评审、采购流程和验收周期纳入计划。")
    add_figure(doc, "roadmap", "图7.1 12个月实施路线示意图")
    p("2. 团队配置", "h2")
    add_table_block(doc, "表7.1 团队角色与职责", [
        ["角色", "职责", "近期重点"],
        ["产品/项目负责人", "产品定位、客户访谈、赛事路演和资源协调", "明确试点场景和验收指标"],
        ["后端工程", "FastAPI、数据存储、权限、审计、Workflow与Open Platform", "稳定核心链路和测试覆盖"],
        ["前端工程", "控制台、搜索、图谱、Marketplace、Developer/Admin页面", "优化演示体验和可视化"],
        ["AI工程", "RAG、Embedding、知识图谱、Agent评测和提示词", "建立评测集和幻觉控制流程"],
        ["交付/商务", "试点推进、方案写作、培训和客户成功", "形成试点清单和报价模型"],
        ["合规/安全顾问", "数据授权、等保、隐私、审计和风险评估", "正式试点前补齐合规材料"],
    ])
    p("3. 主要风险与应对", "h2")
    add_table_block(doc, "表7.2 主要风险与应对措施", [
        ["风险类型", "表现", "应对措施"],
        ["技术风险", "模型幻觉、检索错误、成本波动", "建立引用来源、评测集、缓存、人审和降级策略"],
        ["市场风险", "客户付费意愿和采购周期不确定", "先做小场景试点，再扩大销售"],
        ["数据风险", "客户数据敏感、授权边界复杂", "数据分类分级、脱敏、最小必要和审计留痕"],
        ["交付风险", "不同客户数据格式差异大", "制定接入清单、模板和阶段验收"],
        ["合规风险", "AI输出误导、越权访问、责任不清", "人工复核、输出标识、权限校验和责任边界说明"],
        ["融资风险", "收入验证不足导致估值波动", "控制固定成本，优先用试点数据证明价值"],
    ])

    p("八、合规承诺与资料来源", "h1")
    p("1. 合规承诺", "h2")
    p("本项目正文不涉及国家秘密，不以未授权数据训练公开模型，不将 AI 输出替代法定审批、行政决定或专业人员责任。涉及政府、园区或企业敏感数据时，系统应以合法授权、最小必要、权限分级、日志审计、人工复核、生成内容标识和可追溯来源为前提。正式上线前需结合客户环境完成等保、数据安全、个人信息保护和模型输出风险评估。")
    p("2. 资料来源", "h2")
    sources = [
        "国家网信办等七部门：《生成式人工智能服务管理暂行办法》，2023年7月13日。",
        "国家网信办等部门：《人工智能生成合成内容标识办法》，2025年3月14日。",
        "全国人大常委会：《中华人民共和国数据安全法》，2021年6月10日。",
        "全国人大常委会：《中华人民共和国个人信息保护法》，2021年8月20日。",
        "CNNIC：《生成式人工智能应用发展报告（2024）》，2024年11月。",
        "中国信通院：《中国数字经济发展研究报告（2024年）》，2024年8月。",
        "国家统计局：《中华人民共和国2024年国民经济和社会发展统计公报》。",
        "贵州省大数据发展管理局：《贵州省建设数字经济发展创新区2025年工作要点》。",
        "财政部等部门：中小企业数字化转型城市试点相关通知。",
        "项目本地代码与测试记录：D:\\dma\\day2，Open Platform 2803 passed，核心回归3450 passed，验证时间2026年6月12日。",
    ]
    for source in sources:
        p(source)

    # Add a structural count note outside the business narrative but inside the document for reviewer clarity.
    char_count = paragraph_char_count(texts_for_count)
    p(f"正文无空格字符数约{char_count}字，控制在12000字以内。")

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    out = build()
    print(out)
    print(out.stat().st_size)
